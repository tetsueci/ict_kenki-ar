# -*- coding: utf-8 -*-
"""定格荷重表（.xlsm）→ machines/<機械>.load.json

  並びは terra-maniyac.lsp の read_load_kenki（L7724 付近）と同じ:
    A 列の見出しで行を見つける
      load_boom   … 表の種類
      unitlengh   … 長さの単位（m / mm）
      unitweight  … 重量の単位（t / kN）
      outrigger   … アウトリガのパターン（右へ並ぶ）
      boom        … ブームのパターン（右へ並ぶ）
    そのあと A 列に作業半径の個数、B 列に半径が下へ並ぶ。
    表の値は D 列から。列は [アウトリガ][ブーム] の順に進む。
    空欄は 0（＝その組み合わせは不可）。

  ★機械ごとに表が変わる。機械 1 台につき 1 つ作る。

  使い方:
    python tools/xlsm_to_load.py <表.xlsm> <出力.json>
"""
import io, json, sys
import openpyxl

MARK_COL = 4          # D 列から表の値が始まる


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def main():
    src, out = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src, data_only=True, keep_vba=True)
    ws = wb.worksheets[0]

    # A 列の見出しで行を見つける
    row = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip():
            row.setdefault(v.strip(), r)
    need = ["load_boom", "unitlengh", "unitweight", "outrigger", "boom"]
    for k in need:
        if k not in row:
            raise SystemExit("見出し %s が見つからない（A 列）" % k)

    unit_len = str(ws.cell(row["unitlengh"], MARK_COL).value or "m").strip()
    unit_wt = str(ws.cell(row["unitweight"], MARK_COL).value or "t").strip()
    n_out = int(abs(num(ws.cell(row["outrigger"], MARK_COL - 1).value)))
    n_boom = int(abs(num(ws.cell(row["boom"], MARK_COL - 1).value)))

    outrig = [num(ws.cell(row["outrigger"], MARK_COL + i).value) for i in range(n_out)]
    boom = [num(ws.cell(row["boom"], MARK_COL + i).value) for i in range(n_boom)]

    # 作業半径の個数は、アウトリガ／ブームの下にある数字（A 列）
    r_cnt = None
    for r in range(row["boom"] + 1, ws.max_row + 1):
        v = num(ws.cell(r, 1).value)
        if v:
            r_cnt, r_first = int(abs(v)), r + 1
            break
    if r_cnt is None:
        raise SystemExit("作業半径の個数（A 列の数字）が見つからない")
    radius = [num(ws.cell(r_first + i, 2).value) for i in range(r_cnt)]

    # 表：[アウトリガ][半径][ブーム]
    table = []
    for io_ in range(n_out):
        for ir in range(r_cnt):
            for ib in range(n_boom):
                v = num(ws.cell(r_first + ir, MARK_COL + ib + io_ * n_boom).value)
                table.append(v if v else 0.0)

    d = {
        "note": "定格荷重表。terra-maniyac.lsp の read_load_kenki と同じ並びで読んだ",
        "source": src.replace("\\", "/").split("/")[-1],
        "kind": "rough_terrain_crane",
        "unitLength": unit_len, "unitWeight": unit_wt,
        "outrigger": outrig, "radius": radius, "boom": boom,
        "table": table,
    }
    io.open(out, "w", encoding="utf-8", newline="\n").write(
        json.dumps(d, ensure_ascii=False) + "\n")

    print("アウトリガ %d : %s" % (n_out, outrig))
    print("ブーム     %d : %s" % (n_boom, boom))
    print("作業半径   %d : %s … %s" % (r_cnt, radius[:4], radius[-1]))
    print("表 %d 個（0＝不可 %d 個）/ 単位 %s %s"
          % (len(table), sum(1 for x in table if not x), unit_len, unit_wt))
    print("書いた:", out)


if __name__ == "__main__":
    main()
